from __future__ import annotations

import copy
import hashlib
import json
import re
import shutil
import unicodedata
from collections import Counter, defaultdict
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parent
V3 = ROOT.parent / 'v3' / 'machine'
OUT = ROOT / 'machine'
DOWNLOADS = Path('C:/Users/user/Downloads')
NAMES = [
    'Мааылмат чат ботко 4-5 пункттар боюнча.docx',
    'Простой_запрос_недостающих_данных_AIG2C — копия (1).docx',
    'Условия подачи (1).docx',
    'Простой_запрос_недостающих_данных_AIG2C (3).docx',
    'Чат Бот область (1).docx',
    'Чат Бот боюнча маалымат Ала-Бука РМА.docx',
    'Ала-Бука районунун айыл окмоттору БОТ.docx',
    'Простой_запрос_недостающих_данных_AIG2C — копия.docx',
]
TOPICS = {'воду':'WATER','суу':'WATER','мусор':'WASTE','таштанды':'WASTE',
          'дороги':'ROAD','жолдор':'ROAD','освещение':'STREET_LIGHT','жарыктандыруу':'STREET_LIGHT',
          'люки':'MANHOLE','кудуктар':'MANHOLE','земельные':'LAND'}
PHONE = re.compile(r'(?<!\d)(?:\+?996[\s\u00a0\-]*)?(?:0\d{3}[\s\u00a0\-]*\d{2}[\s\u00a0\-]*\d{2,3}|0\d{9}|\d{3}[\s\u00a0\-]*\d{3}[\s\u00a0\-]*\d{3}|\d{1}[\s\u00a0\-]*\d{2}[\s\u00a0\-]*\d{2})(?!\d)')

def load(name):
    return [json.loads(x) for x in (V3 / (name+'.jsonl')).read_text(encoding='utf8').splitlines() if x.strip()]

def dump(name, rows):
    OUT.mkdir(exist_ok=True)
    (OUT / (name+'.jsonl')).write_text(''.join(json.dumps(x, ensure_ascii=False)+'\n' for x in rows), encoding='utf8')

def key(text):
    text=unicodedata.normalize('NFKC',str(text)).lower().replace('ё','е')
    return re.sub(r'[^\w]+','',text)

def ident(prefix, *parts):
    return prefix+'-'+hashlib.sha256('|'.join(map(str,parts)).encode()).hexdigest()[:14]

def clean(text):
    return ' '.join(str(text).replace('\u00a0',' ').split())

data={p.stem:load(p.stem) for p in V3.glob('*.jsonl')}
data['district_profiles']=[]
data['appeal_flows']=[]
original={n:len(v) for n,v in data.items()}
source_ids={r['source_id'] for r in data['sources']}
org_by_id={o['organization_id']:o for o in data['organizations']}
org_by_name={key(o.get(n)):o for o in data['organizations'] for n in ('name_ru','name_ky') if o.get(n)}
territory_ids={t['territory_id'] for t in data['territories']}
contact_keys={(c['organization_id'],c['type'],key(c['value'])) for c in data['contacts']}
route_keys={(r.get('topic'),r.get('target_organization_id'),r.get('territory_id'),tuple(r.get('source_ids') or [])) for r in data['routing_rules']}
source_additions=defaultdict(Counter)
review=[]
docs={}

def provenance(sid, loc, wording=None):
    p={'source_id':sid,'source_file':docs[sid]['name'],'locator':loc,
       'verification_status':'SOURCE_CONFIRMED','fact_status':'SOURCE_FACT'}
    if wording is not None:p['original_wording']=clean(wording)
    return [p]

def add_territory(kind, name, sid, loc, parent=None):
    name=clean(name)
    tid=f'{kind}:{name}'
    if tid not in territory_ids:
        data['territories'].append({'territory_id':tid,'name_ru':name if kind in ('city','district') else None,
            'name_ky':name if kind not in ('city','district') else None,'type':kind,
            'parent_territory':parent,'parent_status':'SOURCE_CONFIRMED' if parent else 'UNCONFIRMED',
            'source_ids':[sid],'provenance':provenance(sid,loc,name)})
        territory_ids.add(tid);source_additions[sid]['territories']+=1
    return tid

def add_org(name,sid,loc,territory=None,language='ru',aliases=()):
    name=clean(re.sub(r'^\d+[.]\s*','',name))
    norm=key(name)
    existing=org_by_name.get(norm)
    if not existing and 'айыл' in name.lower() and ('окмот' in name.lower() or 'өкмөт' in name.lower()):
        area_names=('Кетмен-Добо','Абды Суеркулов','Элмирбек Иманалиев','Уч-Терек','Толук','Сары-Камыш')
        for area_name in area_names:
            if key(area_name) in norm:
                existing=next((o for o in data['organizations'] if key(area_name) in key(o.get('name_ru') or o.get('name_ky') or '')
                               and ('окмот' in key(o.get('name_ru') or o.get('name_ky') or '') or 'өкмөт' in key(o.get('name_ru') or o.get('name_ky') or ''))),None)
                if existing:break
    if not existing and re.search(r'ДЭП[- №]*\d+',name):
        number=re.search(r'ДЭП[- №]*(\d+)',name).group(1)
        existing=next((o for o in data['organizations'] if re.search(r'(?:ДЭП|предприятие)[- №]*'+number+r'\b',o.get('name_ru') or o.get('name_ky') or '',re.I)
                       and (not territory or not o.get('territory_id') or o.get('territory_id')==territory)),None)
    if not existing and sid=='NEW-265':
        health_matches=[
          ('Центр семейной медицины Сузакского района','Сузак районунун Үй бүлөлүк медицина борбору'),
          ('Территориальная больница Сузакского района','Сузак районунун аймактык ооруканасы'),
          ('Центр общеврачебной практики г.Кок-Жангак','Көк-Жаңгак жалпы дарыгердик практика борбору'),
          ('Центр общеврачебной практики с.Октябрь','Октябрь жалпы дарыгердик практика борбору'),
        ]
        for ru,ky in health_matches:
            if key(ru) in norm:
                existing=next((o for o in data['organizations'] if key(ky) in key(o.get('name_ky') or o.get('name_ru') or '')),None)
                if existing:break
    if not existing:
        canonical={
          key('Токтогульская районная государственная администрация'):'v2-org-9dee194cc6ca',
          key('Ала-Бука райондук мамлекеттик администрациясы'):'v2-org-f79b20b4600b',
          key('Сузакская районная государственная администрация'):'v2-org-7fa33d6d68b9',
          key('Мэрия города Токтогул'):None,
          key('Токтогульское финансовое управление'):'org-kyrgyz-respublikasynyn-finansy-ministrliginin-toktogul-raionduk-f-c446b46',
          key('Учреждение "Сузакское финансовое управление Министерства финансов Кыргызской Республики"'):'org-kyrgyz-respublikasynyn-finansy-ministrliginin-suzak-raionduk-fina-8d50d49',
          key('Сузакское районное ветеринарное управление Ветеринарной службы при Министерстве водных ресурсов, сельского хозяйства и перерабатывающей промышленности Кыргызской Республики'):'v2-org-2ed33d10d37f',
        }
        mapped=canonical.get(norm)
        if mapped:existing=org_by_id.get(mapped)
    if not existing:
        # Reconcile spelling/organizational form only when the distinctive name and territory agree.
        for o in data['organizations']:
            on=key(o.get('name_ru') or o.get('name_ky') or '')
            if len(norm)>13 and (norm in on or on in norm) and min(len(norm),len(on))/max(len(norm),len(on))>.65:
                if not territory or not o.get('territory_id') or o.get('territory_id')==territory:
                    existing=o;break
    if existing:
        oid=existing['organization_id']
        if sid not in existing.get('source_ids',[]):existing.setdefault('source_ids',[]).append(sid)
        alias_field='aliases_ru' if language=='ru' else 'aliases_ky'
        primary_field='name_ru' if language=='ru' else 'name_ky'
        if key(existing.get(primary_field) or '')!=norm and name not in existing.get(alias_field,[]):
            existing.setdefault(alias_field,[]).append(name)
        if norm not in org_by_name:org_by_name[norm]=existing
        return oid
    oid=ident('v4-org',name,territory or '')
    row={'organization_id':oid,'name_ru':name if language=='ru' else None,
         'name_ky':name if language=='ky' else None,'short_name_ru':None,'short_name_ky':None,
         'aliases_ru':list(aliases) if language=='ru' else [],'aliases_ky':list(aliases) if language=='ky' else [],
         'translation_status':'missing','organization_type':'source_named','parent_organization_id':None,
         'territory':[territory] if territory else [],'territory_id':territory,
         'territory_status':'SOURCE_CONTEXT' if territory else 'NEEDS_CONFIRMATION','status':'SOURCE_CONFIRMED',
         'source_ids':[sid],'provenance':provenance(sid,loc,name),'verification_status':'SOURCE_CONFIRMED',
         'visibility':'PUBLIC'}
    data['organizations'].append(row);org_by_name[norm]=row;org_by_id[oid]=row
    source_additions[sid]['organizations']+=1
    return oid

def add_contact(oid,ctype,value,sid,loc,territory=None,visibility='PUBLIC',status='SOURCE_CONFIRMED',purpose='office_contact',wording=None):
    value=clean(value)
    if not value:return
    k=(oid,ctype,key(value))
    def phone_key(v):
        digits=re.sub(r'\D','',v)
        if digits.startswith('996'):digits=digits[3:]
        return digits.lstrip('0')
    for old in data['contacts']:
        if old.get('organization_id')!=oid or old.get('type')!=ctype:continue
        same=key(old.get('value'))==key(value)
        if ctype=='phone':same=same or (phone_key(old.get('value',''))==phone_key(value) and len(phone_key(value))>=7)
        if same:
            if sid not in old.get('source_ids',[]):
                old.setdefault('source_ids',[]).append(sid)
                old.setdefault('provenance',[]).extend(provenance(sid,loc,wording or value))
            return old['contact_id']
    cid=ident('v4-contact',oid,ctype,value,sid)
    row={'contact_id':cid,'organization_id':oid,'type':ctype,'contact_type':ctype,'value':value,
         'purpose':purpose,'territory_id':territory,'source_ids':[sid],
         'provenance':provenance(sid,loc,wording or value),'verification_status':'SOURCE_CONFIRMED',
         'pilot_status':status,'status':status,'visibility':visibility,
         'publication_status':'PILOT_PUBLIC' if visibility=='PUBLIC' and status=='SOURCE_CONFIRMED' else 'INTERNAL_ONLY'}
    data['contacts'].append(row);contact_keys.add(k);source_additions[sid]['contacts']+=1
    if ctype in ('working_hours','reception_hours'):data['hours'].append(copy.deepcopy(row))
    return cid

def add_comp(oid,text,sid,loc,territory=None,lang='ru',topic='OTHER'):
    text=clean(text)
    if not text:return
    cid=ident('v4-comp',oid,text,sid)
    data['competences'].append({'competence_id':cid,'organization_id':oid,'topic':topic,
        'description_ru':text if lang=='ru' else None,'description_ky':text if lang=='ky' else None,
        'territory_id':territory,'source_ids':[sid],'provenance':provenance(sid,loc,text),
        'verification_status':'SOURCE_CONFIRMED','status':'SOURCE_CONFIRMED','visibility':'PUBLIC'})
    source_additions[sid]['competences']+=1
    add_chunk(oid,text,sid,loc,territory,lang,topic,['competence'])

def add_chunk(oid,text,sid,loc,territory=None,lang='ru',topic='OTHER',tags=None,route_id=None):
    text=clean(text)
    if not text:return
    row={'chunk_id':ident('v4-chunk',sid,loc,oid,topic,territory or '',text[:60]),'document_id':sid,'source_id':sid,
         'organization_ids':[oid] if oid else [],'topic':topic,'territory_scope':territory,
         'language':lang,'title':(org_by_id.get(oid) or {}).get('name_ru') or (org_by_id.get(oid) or {}).get('name_ky') or 'Источник',
         'text':text,'source_file':docs[sid]['name'],'locator':loc,'tags':tags or [],
         'routing_rule_id':route_id,'provenance':provenance(sid,loc,text),
         'verification_status':'SOURCE_CONFIRMED','publication_status':'PILOT_PUBLIC'}
    data['rag_chunks'].append(row);source_additions[sid]['rag_chunks']+=1

def add_route(topic,oid,sid,loc,territory,wording,secondary=None,lang='ru',zones=None):
    if not territory:
        review.append(f'{sid} {loc}: маршрут без территории: {wording}')
        return
    rid=ident('v4-route',topic,oid,territory,sid,loc)
    data['routing_rules'].append({'routing_rule_id':rid,'topic':topic,
      'canonical_topic':topic,'original_wording':clean(wording),'trigger_phrases_ru':[clean(wording)] if lang=='ru' else [],
      'trigger_phrases_ky':[clean(wording)] if lang=='ky' else [],'target_organization_id':oid,
      'secondary_organization_ids':secondary or [],'multi_agency':bool(secondary),
      'route_type':'SOURCE_CONFIRMED_SUGGESTION','operator_confirmation_required':True,
      'territory_id':territory,'territory_scope':territory,'service_zones':zones or [],
      'source_ids':[sid],'provenance':provenance(sid,loc,wording),'status':'SOURCE_CONFIRMED',
      'visibility':'PUBLIC'})
    source_additions[sid]['routing_rules']+=1
    if oid:add_chunk(oid,f'{topic}: {org_by_id[oid].get("name_ru") or org_by_id[oid].get("name_ky")}',sid,loc,territory,lang,topic,['routing'],rid)

def add_zone(oid,tid,sid,loc,wording):
    rid=ident('v4-zone',oid,tid,sid,loc)
    if any(x.get('relation_id')==rid for x in data['service_territories']):return
    data['service_territories'].append({'relation_id':rid,'organization_id':oid,'territory_id':tid,
      'original_wording':clean(wording),'status':'SOURCE_CONFIRMED','source_ids':[sid],
      'provenance':provenance(sid,loc,wording)})
    source_additions[sid]['service_territories']+=1

def add_staff(oid,name,position,phone,sid,loc,territory=None):
    name=clean(name);phone=clean(phone)
    if not name and not phone:return
    staff_role_id=ident('v4-staff',oid,name,phone,sid,loc)
    if any(s.get('staff_role_id')==staff_role_id for s in data['staff_roles']):return
    data['staff_roles'].append({'staff_role_id':staff_role_id,'staff_id':ident('staff',oid,name),
      'organization_id':oid,'person_name':name,'position':clean(position),'phone':phone or None,
      'system_role':'manager' if any(x in position.lower() for x in ('руковод','директор','жетекчи','мэр','аким')) else 'executor',
      'visibility':'INTERNAL_ONLY','territory_id':territory,'source_ids':[sid],
      'provenance':provenance(sid,loc,name+' '+position+' '+phone),'status':'SOURCE_CONFIRMED'})
    source_additions[sid]['staff_roles']+=1

def add_profile(territory,oid,population,counts,sid,loc):
    data['district_profiles'].append({'profile_id':ident('v4-profile',sid,territory),
      'territory_id':territory,'organization_id':oid,'population':population,'facility_counts':counts,
      'source_ids':[sid],'provenance':provenance(sid,loc),
      'status':'SOURCE_CONFIRMED','visibility':'PUBLIC'})
    source_additions[sid]['district_profiles']+=1

for i,name in enumerate(NAMES,258):
    path=DOWNLOADS/name
    raw=path.read_bytes();sid=f'NEW-{i}';d=Document(path)
    docs[sid]={'name':name,'doc':d,'hash':hashlib.sha256(raw).hexdigest()}
    relation_to={258:'NEW-259',259:'NEW-261',264:'NEW-263',265:'NEW-261'}
    relation='SUPPLEMENTS' if i in relation_to else 'NEW_SOURCE'
    data['sources'].append({'source_id':sid,'file_name':name,'source_file':name,'file_type':'docx',
       'size':len(raw),'sha256':docs[sid]['hash'],'language':'ky' if i in (262,263,264) else 'ru',
       'source_relation':relation,'relation_to':relation_to.get(i),
       'status':'processed','normalization_status':'promoted','publication_status':'PILOT_PUBLIC',
       'paragraph_count':len(d.paragraphs),'table_count':len(d.tables)})

# Filled Toktogul district registry and official contacts (NEW-259).
sid='NEW-259';d=docs[sid]['doc'];district=add_territory('district','Токтогул',sid,'paragraph:3')
tok_rga=add_org('Токтогульская районная государственная администрация',sid,'paragraph:5',district)
for ctype,val,loc in [('address',d.paragraphs[7].text,'paragraph:7'),('phone',d.paragraphs[9].text,'paragraph:9'),
 ('email',d.paragraphs[16].text,'paragraph:16'),('reception_hours','Пн, Ср, Чт, Пт 14:00–17:00', 'paragraph:10-14')]:
    add_contact(tok_rga,ctype,val,sid,loc,district)
add_profile(district,tok_rga,108938,{'schools':41,'kindergartens':26,'hospitals':1,
 'family_medical_centers':15,'culture_clubs':5,'libraries':24,'sports':95},sid,'paragraph:17-32')
data['escalation_rules'].append({'escalation_rule_id':ident('v4-escalation',sid,'paragraph:63'),
 'organization_id':tok_rga,'territory_id':district,'rule_type':'REFER_TO_COMPETENT_AUTHORITY',
 'description':clean(d.paragraphs[63].text),'multi_agency':False,'production_status':'NEEDS_OPERATOR',
 'source_ids':[sid],'provenance':provenance(sid,'paragraph:63',d.paragraphs[63].text),'visibility':'PUBLIC'})
source_additions[sid]['escalation_rules']+=1
data['procedures'].append({'procedure_id':ident('v4-proc',sid,'paragraph:63'),
 'organization_id':tok_rga,'topic':'APPEAL','name':'Передача обращения компетентному органу',
 'steps':[clean(d.paragraphs[63].text)],'required_documents':[],'conditions':'Вопрос не относится к компетенции службы',
 'deadline':None,'territory_id':district,'completeness':'PARTIAL','source_ids':[sid],
 'provenance':provenance(sid,'paragraph:63',d.paragraphs[63].text),'status':'SOURCE_CONFIRMED','visibility':'PUBLIC'})
source_additions[sid]['procedures']+=1
for ri,row in enumerate(d.tables[1].rows[1:],1):
    cells=[clean(c.text) for c in row.cells]
    name=re.sub(r'^\d+[.]\s*','',cells[0]);phone=cells[1];desc=cells[2]
    if not name:continue
    oid=add_org(name,sid,f'table:1:row:{ri}',district)
    for num in re.findall(r'(?<!\d)(?:\+?996)?0?\d[\d\s\u00a0\-]{5,13}\d(?!\d)',phone):
        num=clean(num)
        if len(re.sub(r'\D','',num))>=7:add_contact(oid,'phone',num,sid,f'table:1:row:{ri}',district,wording=' | '.join(cells))
    if desc:add_comp(oid,desc,sid,f'table:1:row:{ri}',district)
for pi,topic in [(47,'WATER'),(48,'MANHOLE'),(49,'ROAD'),(50,'STREET_LIGHT'),(51,'WASTE'),(52,'LAND'),(53,'ILLEGAL_CONSTRUCTION'),(54,'ELECTRICITY'),(55,'EDUCATION'),(56,'HEALTH'),(57,'SOCIAL_SUPPORT'),(59,'PUBLIC_ORDER')]:
    txt=clean(d.paragraphs[pi].text)
    add_chunk(tok_rga,txt,sid,f'paragraph:{pi}',district,'ru',topic,['routing_context'])
res=add_org('Токтогул РЭС',sid,'paragraph:65',district)
add_zone(res,district,sid,'paragraph:65',d.paragraphs[65].text)
add_route('ELECTRICITY',res,sid,'paragraph:54',district,d.paragraphs[54].text)
dep23=add_org('ДЭП №23',sid,'paragraph:66',district)
for area in ('Кетмен-Добо','Абды Суеркулов'):
 tid=add_territory('ayyl_area',area,sid,'paragraph:66',district)
 add_zone(dep23,tid,sid,'paragraph:66',d.paragraphs[66].text)
 add_route('ROAD',dep23,sid,'paragraph:66',tid,d.paragraphs[66].text)

# Toktogul ayyl areas and municipality (NEW-258).
sid='NEW-258';d=docs[sid]['doc']
area_specs=[('Кетмен-Добо',1,2,3,4,5,11),('Толук',13,15,16,17,18,24),
 ('Абды Суеркулов',25,27,28,29,30,36),('Элмирбек Иманалиев',40,42,43,46,47,53),
 ('Уч-Терек',56,58,59,60,61,68),('Сары-Камыш',69,70,71,72,73,80)]
for area,heading,orgp,addrp,villp,route_start,route_end in area_specs:
    tid=add_territory('ayyl_area',area,sid,f'paragraph:{heading}',district)
    line=clean(d.paragraphs[orgp].text)
    name='Айыл окмоту '+area+' айылного аймака'
    oid=add_org(name,sid,f'paragraph:{orgp}',tid)
    nums=re.findall(r'\+996[\s\u00a0\-]*\d[\d\s\u00a0\-]{5,13}\d|0\d{3}[\s\u00a0\-]*\d{3}[\s\u00a0\-]*\d{3}',line)
    for num in nums:
        if len(re.sub(r'\D','',num))<10:
            add_contact(oid,'phone',num,sid,f'paragraph:{orgp}',tid,'INTERNAL_ONLY','AMBIGUOUS',wording=line)
            review.append(f'{sid} paragraph:{orgp}: неполный телефон {num}')
        else:add_contact(oid,'phone',num,sid,f'paragraph:{orgp}',tid,wording=line)
    addr=clean(d.paragraphs[addrp].text)
    matches=list(re.finditer(r'график работы',addr,re.I))
    if matches:
        last=matches[-1];before=addr[:last.start()];hours=addr[last.end():].strip(' .,:–-')
    else:before=addr;hours=None
    before=re.sub(r'^Адрес(?:, общий телефон, график работы)?[. :–-]*','',before,flags=re.I).strip(' .,:–-')
    if before:add_contact(oid,'address',before,sid,f'paragraph:{addrp}',tid)
    if hours:add_contact(oid,'working_hours',hours,sid,f'paragraph:{addrp}',tid)
    village_line=clean(d.paragraphs[villp].text)
    if ':' in village_line:village_line=village_line.split(':',1)[1]
    if 'села;' in village_line:village_line=village_line.split('села;',1)[1]
    for village in village_line.split(','):
        village=clean(village).strip(' .')
        if village and len(village)<45:add_territory('village',village,sid,f'paragraph:{villp}',tid)
    for pi in range(route_start,route_end):
        txt=clean(d.paragraphs[pi].text)
        for token,topic in TOPICS.items():
            if token in txt.lower() and ('отвечает' in txt.lower() or 'направлять' in txt.lower()):
                mp={'Кетмен-Добо':'МП «Таза аймак»','Абды Суеркулов':'МП «Торкен»',
                    'Элмирбек Иманалиев':'МП «Жаңы Доор»','Уч-Терек':'МП «Уч-Терек таза»'}
                use_mp=area in mp and topic in ('WATER','WASTE','MANHOLE')
                if area=='Абды Суеркулов' and topic in ('WATER','MANHOLE'):use_mp=False
                if area=='Сары-Камыш' and topic=='ROAD':
                    dep=add_org('ДЭП №23',sid,f'paragraph:{pi}',tid)
                    add_route(topic,oid,sid,f'paragraph:{pi}',tid,txt,secondary=[dep]);break
                if area=='Сары-Камыш' and topic=='STREET_LIGHT':
                    review.append(f'{sid} paragraph:{pi}: РЭС не указан точно; маршрут к айыл өкмөтү с подтверждением оператора')
                target_id=add_org(mp[area],sid,f'paragraph:{pi}',tid) if use_mp else oid
                add_route(topic,target_id,sid,f'paragraph:{pi}',tid,txt)
                break
city=add_territory('city','Токтогул',sid,'table:0:row:1',district)
mayor=add_org('Мэрия города Токтогул',sid,'table:0:row:1',city)
cell=clean(d.tables[0].rows[1].cells[2].text)
add_contact(mayor,'address','г. Токтогул ул Черноткач №5',sid,'table:0:row:1',city,wording=cell)
add_contact(mayor,'phone','5-02-78',sid,'table:0:row:1',city,wording=cell)
add_contact(mayor,'reception_hours','Каждый вторник с 14:00-16:00 мэр города Каждую пятницу с 10:00-12:00 мэр города',
            sid,'table:0:row:1',city,wording=cell)
for line in d.tables[0].rows[2].cells[2].text.splitlines():
    line=clean(line)
    if '«Суу канал»' in line:
        oid=add_org('МП «Суу канал»',sid,'table:0:row:2',city)
        for topic in ('WATER','SEWER'):add_route(topic,oid,sid,'table:0:row:2',city,line)
    if '«Ыңгайлуу шарт»' in line:
        oid=add_org('МП «Ыңгайлуу шарт»',sid,'table:0:row:2',city)
        for topic in ('ROAD','STREET_LIGHT','WASTE'):add_route(topic,oid,sid,'table:0:row:2',city,line)
for ri,row in enumerate(d.tables[1].rows[1:],1):
    school=clean(row.cells[1].text);kindergarten=clean(row.cells[2].text)
    if school:add_org('Школа '+school,sid,f'table:1:row:{ri}',city)
    if kindergarten:add_org('Детский сад '+kindergarten,sid,f'table:1:row:{ri}',city)

# Tash-Komur municipality (NEW-260).
sid='NEW-260';d=docs[sid]['doc'];city=add_territory('city','Таш-Көмүр',sid,'paragraph:0')
mayor=add_org('Таш-Көмүр шаардык мэриясы',sid,'table:0:row:1',city,'ky')
for ctype,ri in [('address',5),('working_hours',6),('reception_hours',7)]:
 add_contact(mayor,ctype,d.tables[0].rows[ri].cells[1].text,sid,f'table:0:row:{ri}',city)
add_contact(mayor,'phone','03745 50030',sid,'paragraph:0',city,'INTERNAL_ONLY','CONFLICT')
add_contact(mayor,'phone','0745 50030',sid,'table:0:row:4',city,'INTERNAL_ONLY','CONFLICT')
data['conflicts'].append({'conflict_id':ident('v4-conflict',sid,'mayor-phone'),'entity':'Таш-Көмүр шаардык мэриясы',
 'organization_id':mayor,'field':'phone','values':['03745 50030','0745 50030'],
 'source_ids':[sid],'provenance':[{'source_id':sid,'source_file':docs[sid]['name'],'locator':'paragraph:0'},
                      {'source_id':sid,'source_file':docs[sid]['name'],'locator':'table:0:row:4'}],
 'status':'CONFLICT','resolution':'NO_WINNER','visibility':'INTERNAL_ONLY'})
source_additions[sid]['conflicts']+=1
add_comp(mayor,d.tables[0].rows[3].cells[1].text,sid,'table:0:row:3',city)
for pi,topic,n in [(2,'WATER','Х/З предприятие «Суу чарба»'),(3,'ROAD','ДЭУ №38'),
                   (4,'STREET_LIGHT','Департамент муниципальной собственности Таш-Көмүр'),
                   (5,'WASTE','МП «Таш-Көмүр-Тазалык»')]:
 oid=add_org(n,sid,f'paragraph:{pi}',city);add_route(topic,oid,sid,f'paragraph:{pi}',city,d.paragraphs[pi].text)
for zone in ('Кызыл-Алма','Кызыл-Жар','Мендигулов','Тынчтык','Түндүк'):
    zone_id=add_territory('service_zone',zone,sid,'table:0:row:10',city)
    add_zone(mayor,zone_id,sid,'table:0:row:10',d.tables[0].rows[10].cells[1].text)
edu=add_org('Городской отдел образования Таш-Көмүр',sid,'paragraph:9',city)
for num in ('03745 50207','0772 785 070'):add_contact(edu,'phone',num,sid,'paragraph:9',city,wording=d.paragraphs[9].text)
data['procedures'].append({'procedure_id':ident('v4-proc',sid,'table:0:row:8'),'organization_id':mayor,
 'topic':'APPEAL','name':'Подача обращения в мэрию Таш-Көмүр','steps':[clean(d.tables[0].rows[8].cells[1].text)],
 'required_documents':[],'conditions':None,'deadline':None,'territory_id':city,'completeness':'PARTIAL',
 'source_ids':[sid],'provenance':provenance(sid,'table:0:row:8',d.tables[0].rows[8].cells[1].text),
 'status':'SOURCE_CONFIRMED','visibility':'PUBLIC'});source_additions[sid]['procedures']+=1
add_chunk(mayor,d.tables[0].rows[8].cells[1].text,sid,'table:0:row:8',city,'ru','APPEAL',['procedure'])
data['appeal_flows'].append({'flow_id':ident('v4-flow',sid,'table:0:row:8'),'organization_id':mayor,
 'territory_id':city,'steps':['Подать обращение через приёмную или QR-код у главного входа'],
 'source_ids':[sid],'provenance':provenance(sid,'table:0:row:8',d.tables[0].rows[8].cells[1].text),
 'status':'SOURCE_CONFIRMED','visibility':'PUBLIC'})
source_additions[sid]['appeal_flows']+=1

# Mayluu-Suu filled card in the updated request form (NEW-261).
sid='NEW-261';d=docs[sid]['doc'];city=add_territory('city','Майлуу-Суу',sid,'table:4:row:1')
mayor=add_org('Мэрия города Майлуу-Суу',sid,'table:4:row:1',city)
for ri,row in enumerate(d.tables[4].rows[1:],1):
 q=clean(row.cells[0].text).lower();answer=clean(row.cells[1].text)
 if not answer or answer in ('-', 'нет'):continue
 ctype=next((t for needle,t in [('публичный телефон','phone'),('адрес приёма','address'),('график работы','working_hours'),('часы личного','reception_hours')] if needle in q),None)
 if ctype:add_contact(mayor,ctype,answer,sid,f'table:4:row:{ri}',city,wording=q+' | '+answer)
 elif 'отвечает' in q:add_comp(mayor,answer,sid,f'table:4:row:{ri}',city)

# Regional health institution table (NEW-262): named staff phones remain internal.
sid='NEW-262';d=docs[sid]['doc']
for ri,row in enumerate(d.tables[0].rows[1:],1):
 cells=[clean(c.text) for c in row.cells]
 if len(cells)<6 or not cells[1]:continue
 name,address,competence,services,staff=cells[1:6]
 territory=None
 m=re.search(r'([А-Яа-яӨҮөүЁё-]+) (шаары|району|айылы)',address)
 if m:territory=add_territory('city' if m.group(2)=='шаары' else 'district' if m.group(2)=='району' else 'village',m.group(1),sid,f'table:0:row:{ri}')
 oid=add_org(name,sid,f'table:0:row:{ri}',territory,'ky')
 if address:add_contact(oid,'address',address,sid,f'table:0:row:{ri}',territory,wording=' | '.join(cells[:3]))
 if competence:add_comp(oid,competence,sid,f'table:0:row:{ri}',territory,'ky','HEALTH')
 if services:
  data['services'].append({'service_id':ident('v4-service',oid,sid,ri),'organization_id':oid,'name':'Медициналык кайрылуулар',
   'description':services,'topic':'HEALTH','territory_id':territory,'source_ids':[sid],
   'provenance':provenance(sid,f'table:0:row:{ri}',services),'status':'SOURCE_CONFIRMED','visibility':'PUBLIC'})
  source_additions[sid]['services']+=1
  add_chunk(oid,services[:850],sid,f'table:0:row:{ri}',territory,'ky','HEALTH',['service'])
 for line in staff.replace(';','\n').split(' / '):
  nums=re.findall(r'(?:\+?996[\s\u00a0\-]*)?0\d{9}|\+996[\d\s\u00a0\-]{9,16}',line)
  if nums:
   first=re.split(r'тел\.?\s*:?|\+?996|0\d{9}',line,flags=re.I)[0]
   for num in nums:add_staff(oid,first,'Сотрудник/ответственный',num,sid,f'table:0:row:{ri}',territory)

# Ala-Buka district administration (NEW-263).
sid='NEW-263';d=docs[sid]['doc'];district=add_territory('district','Ала-Бука',sid,'paragraph:0')
rga=add_org('Ала-Бука райондук мамлекеттик администрациясы',sid,'paragraph:17',district,'ky')
add_contact(rga,'address',d.paragraphs[1].text.split(':',1)[1],sid,'paragraph:1',district)
for pi in (2,3):add_contact(rga,'phone',d.paragraphs[pi].text.split(':')[-1],sid,f'paragraph:{pi}',district)
add_contact(rga,'reception_hours',d.paragraphs[4].text.split(':',1)[1],sid,'paragraph:4',district)
add_contact(rga,'email',d.paragraphs[5].text.split(':',1)[1],sid,'paragraph:5',district)
add_staff(rga,'Абыкеев Темирболот Асаналиевич','аким','0990 907050',sid,'paragraph:17-18',district)
add_profile(district,rga,122042,{'schools':42,'kindergartens':34,'hospitals':1,
 'medical_facilities':34,'culture_clubs':16,'libraries':22,'sports':42},sid,'paragraph:6-15')

# Ala-Buka ayyl areas, municipal enterprises and explicit route rows (NEW-264).
sid='NEW-264';d=docs[sid]['doc']
for ti,table in enumerate(d.tables):
 first=clean(table.rows[0].cells[1].text);name=first.split('\n')[0]
 area=['Айры-Там','Ала-Бука','Т.Балтагулов','Ынтымак'][ti]
 tid=add_territory('ayyl_area',area,sid,f'table:{ti}:row:0',district)
 if ti==1:oid=add_org('Ала-Бука айыл өкмөтү',sid,f'table:{ti}:row:0',tid,'ky')
 else:oid=add_org(re.split(r'\+996|\b0\d{3}\b',first)[0].strip(),sid,f'table:{ti}:row:0',tid)
 for num in re.findall(r'\+996\s*\d{9}|0\d{9}|0\d{4}\s*\d{5}',first):add_contact(oid,'phone',num,sid,f'table:{ti}:row:0',tid,wording=first)
 addr=clean(table.rows[1].cells[1].text)
 addr_only=addr.split('ПН-ПТ')[0].split('жумуш убактысы')[0].split('Тел.')[0]
 addr_only=re.split(r'\+996\s*\d|0\d{3}\s*\d',addr_only)[0]
 add_contact(oid,'address',addr_only,sid,f'table:{ti}:row:1',tid)
 hm=re.search(r'(?:ПН-ПТ|жумуш убактысы:).*',addr,re.I)
 if hm:add_contact(oid,'working_hours',hm.group(0),sid,f'table:{ti}:row:1',tid)
 for num in re.findall(r'\+996\s*\d{9}|0\d{9}|\b\d-\d{2}-\d{2}\b',addr):
  add_contact(oid,'phone',num,sid,f'table:{ti}:row:1',tid,wording=addr)
 villages=clean(table.rows[2].cells[1].text)
 for village in re.split(r'[,\n]',villages):
  village=clean(village)
  if village and len(village)<50:add_territory('village',village,sid,f'table:{ti}:row:2',tid)
 for ri,row in enumerate(table.rows[3:],3):
  q=clean(row.cells[0].text);ans=clean(row.cells[1].text)
  if not ans:continue
  topic=next((v for k,v in TOPICS.items() if k in q.lower()),None)
  if topic:
   mp_by_area={0:'МИ «1-Май Эмгек»',1:{3:'МП «Ала-Бука Водоканал»',4:'МП «Ала-Бука Тазалык»',
             5:'МП «Ала-Бука Тазалык»',6:'МП «Ала-Бука Тазалык»',7:'МП «Ала-Бука Водоканал»'},
             2:'МИ «Чанач-Сай»',3:'МИ «Ынтымак Тазалык»'}
   target=oid
   if ti==2 and ri==7:
    target=None;review.append(f'{sid} table:{ti}:row:{ri}: указан только специалист по ЧС; организация люков не определена')
   elif topic=='LAND':target=oid
   elif ti==3 and ri in (3,5,7):target=oid
   elif ti in mp_by_area:
    mpname=mp_by_area[ti]
    if isinstance(mpname,dict):mpname=mpname.get(ri)
    if mpname:target=add_org(mpname,sid,f'table:{ti}:row:{ri}',tid)
   if target:
    second=[]
    if ti==1 and ri==5:second=[add_org('МП «Ала-Бука Техсервис»',sid,f'table:{ti}:row:{ri}',tid)]
    add_route(topic,target,sid,f'table:{ti}:row:{ri}',tid,q+' '+ans,secondary=second,lang='ky' if ti==3 else 'ru')
  if 'Если вопрос не их' in q or 'компетенциясына кирбеген' in q:
   data['escalation_rules'].append({'escalation_rule_id':ident('v4-escalation',sid,ti,ri),
     'organization_id':oid,'territory_id':tid,'rule_type':'REFER_TO_COMPETENT_AUTHORITY',
     'description':ans,'multi_agency':'же' in ans or 'или' in ans,'production_status':'NEEDS_OPERATOR',
     'source_ids':[sid],'provenance':provenance(sid,f'table:{ti}:row:{ri}',q+' '+ans),'visibility':'PUBLIC'})
   source_additions[sid]['escalation_rules']+=1
  # Named individuals' numbers are internal, including enterprise managers.
  nums=re.findall(r'\+996\s*\d{9}|0\d{9}|(?<!\d)\d{3}\s*\d{3}\s*\d{3}(?!\d)',ans)
  if nums:
   target=target if topic and target else oid
   for num in nums:add_staff(target,ans.split('тел')[0][:90],'Ответственный по вопросу',num,sid,f'table:{ti}:row:{ri}',tid)

# Suzak filled municipal registry in NEW-265 (table 4); named MП leaders in table 3 stay private.
sid='NEW-265';d=docs[sid]['doc'];district=add_territory('district','Сузак',sid,'table:4:row:1')
suzak_rga=add_org('Сузакская районная государственная администрация',sid,'paragraph:21',district)
add_contact(suzak_rga,'address',d.paragraphs[23].text,sid,'paragraph:23',district)
new_suzak_phone=clean(d.paragraphs[25].text)
def local_digits(value):
    number=re.sub(r'\D','',value)
    return number[3:] if number.startswith('996') else number.lstrip('0')
old_faxes=[c for c in data['contacts'] if c.get('organization_id')==suzak_rga and c.get('type')=='fax'
            and local_digits(c.get('value',''))==local_digits(new_suzak_phone)]
if old_faxes:
    for old_fax in old_faxes:
        old_fax['visibility']='INTERNAL_ONLY';old_fax['status']='CONFLICT';old_fax['publication_status']='INTERNAL_ONLY'
    add_contact(suzak_rga,'phone',new_suzak_phone,sid,'paragraph:25',district,'INTERNAL_ONLY','CONFLICT')
    data['conflicts'].append({'conflict_id':ident('v4-conflict',sid,'suzak-phone-fax'),
      'entity':'Сузак райондук мамлекеттик администрациясы','organization_id':suzak_rga,
      'field':'contact_type','values':['fax: '+old_faxes[0]['value'],'phone: '+new_suzak_phone],
      'source_ids':[old_faxes[0]['source_ids'][0],sid],
      'provenance':old_faxes[0].get('provenance',[])+provenance(sid,'paragraph:25',new_suzak_phone),
      'status':'CONFLICT','resolution':'NO_WINNER','visibility':'INTERNAL_ONLY'})
    source_additions[sid]['conflicts']+=1
else:add_contact(suzak_rga,'phone',new_suzak_phone,sid,'paragraph:25',district)
add_contact(suzak_rga,'email',d.paragraphs[30].text,sid,'paragraph:30',district)
add_contact(suzak_rga,'reception_hours','Понедельник 14:00–15:00; Пятница 09:00–10:00',sid,'paragraph:26-28',district)
add_staff(suzak_rga,'Кулубаев Дамир Садырбекович','аким',d.paragraphs[54].text.split(':')[-1],sid,'paragraph:52-54',district)
add_profile(district,suzak_rga,287641,{'schools':109,'kindergartens':74,'hospitals':1,
 'general_practice_centers':2,'medical_facilities':74,'culture_clubs':17,'libraries':33,'sports':77},sid,'paragraph:31-46')
data['escalation_rules'].append({'escalation_rule_id':ident('v4-escalation',sid,'paragraph:79'),
 'organization_id':suzak_rga,'territory_id':district,'rule_type':'REFER_TO_COMPETENT_AUTHORITY',
 'description':clean(d.paragraphs[79].text),'multi_agency':False,'production_status':'NEEDS_OPERATOR',
 'source_ids':[sid],'provenance':provenance(sid,'paragraph:79',d.paragraphs[79].text),'visibility':'PUBLIC'})
source_additions[sid]['escalation_rules']+=1
data['procedures'].append({'procedure_id':ident('v4-proc',sid,'paragraph:79'),
 'organization_id':suzak_rga,'topic':'APPEAL','name':'Передача обращения компетентному органу',
 'steps':[clean(d.paragraphs[79].text)],'required_documents':[],'conditions':'Вопрос не относится к компетенции службы',
 'deadline':None,'territory_id':district,'completeness':'PARTIAL','source_ids':[sid],
 'provenance':provenance(sid,'paragraph:79',d.paragraphs[79].text),'status':'SOURCE_CONFIRMED','visibility':'PUBLIC'})
source_additions[sid]['procedures']+=1
for ri,row in enumerate(d.tables[4].rows[1:],1):
 cells=[clean(c.text) for c in row.cells]
 if len(cells)<3 or not cells[0]:continue
 name=re.sub(r'^\d+[.]\s*','',cells[0]);oid=add_org(name,sid,f'table:4:row:{ri}',district)
 for num in re.findall(r'(?<!\d)0\d{9}(?!\d)',cells[1]):
  add_contact(oid,'phone',num,sid,f'table:4:row:{ri}',district,wording=' | '.join(cells))
 if cells[2]:add_comp(oid,cells[2],sid,f'table:4:row:{ri}',district)
for ri,row in enumerate(d.tables[3].rows[1:],1):
 cells=[clean(c.text) for c in row.cells]
 if len(cells)<4 or not cells[1]:continue
 oid=add_org(cells[1],sid,f'table:3:row:{ri}',district)
 if cells[2] or cells[3]:add_staff(oid,cells[2],'Руководитель',cells[3],sid,f'table:3:row:{ri}',district)
service_zone_specs={81:('Сузак РЭС',['Кара-Дарьин','Сузак','Барпы','Сайып-Атабеков','Кызыл-Туу'],'ELECTRICITY'),
 82:('Сузакское предприятие электрических сетей',['Кара-Дарьин','Сузак','Барпы','Сайып-Атабеков','Кызыл-Туу'],'ELECTRICITY'),
 83:('Кок-Арт РЭС',['Көк-Арт','Багыш','Кыз-Көл'],'ELECTRICITY'),
 84:('ДЭП-31',['Багыш','Көк-Арт'],'ROAD'),85:('ДЭП-52',['Кызыл-Туу','Кыз-Көл'],'ROAD'),
 86:('ДЭП-22',['Сузак','Сайып-Атабеков','Барпы','Көк-Арт'],'ROAD')}
for pi,(name,areas,topic) in service_zone_specs.items():
 oid=add_org(name,sid,f'paragraph:{pi}',district)
 for area in areas:
  tid=add_territory('ayyl_area',area,sid,f'paragraph:{pi}',district)
  add_zone(oid,tid,sid,f'paragraph:{pi}',d.paragraphs[pi].text)
  if pi==82:
   for prior in data['routing_rules']:
    if prior.get('topic')=='ELECTRICITY' and prior.get('territory_scope')==tid and sid in prior.get('source_ids',[]):
     prior.setdefault('secondary_organization_ids',[]).append(oid)
     prior['multi_agency']=True
   continue
  if topic=='ROAD' and area=='Көк-Арт':continue
  add_route(topic,oid,sid,f'paragraph:{pi}',tid,d.paragraphs[pi].text)
 if pi==83:
  city_id=add_territory('city','Көк-Жаңгак',sid,'paragraph:83',district)
  add_zone(oid,city_id,sid,'paragraph:83',d.paragraphs[83].text)
  add_route('ELECTRICITY',oid,sid,'paragraph:83',city_id,d.paragraphs[83].text)
data['conflicts'].append({'conflict_id':ident('v4-conflict',sid,'kok-art-road'),
 'entity':'Көк-Арт айыл аймагы','field':'ROAD routing','values':['ДЭП-31','ДЭП-22'],
 'source_ids':[sid],'provenance':[{'source_id':sid,'source_file':docs[sid]['name'],'locator':'paragraph:84'},
   {'source_id':sid,'source_file':docs[sid]['name'],'locator':'paragraph:86'}],
 'status':'CONFLICT','resolution':'NO_WINNER','visibility':'INTERNAL_ONLY'})
source_additions[sid]['conflicts']+=1

# Preserve the wording of new source categories as lexical evidence.
for sid,paragraphs,tid in [
    ('NEW-259',range(47,60), 'district:Токтогул'),
    ('NEW-265',range(63,77), 'district:Сузак')]:
    doc=docs[sid]['doc']
    for pi in paragraphs:
        wording=clean(doc.paragraphs[pi].text)
        if not wording:continue
        topic=next((v for k,v in TOPICS.items() if k in wording.lower()),'OTHER')
        data['appeal_categories'].append({'category_id':ident('v4-category',sid,pi),
          'original_wording':wording,'canonical_topic':topic,'territory_id':tid,
          'source_ids':[sid],'provenance':provenance(sid,f'paragraph:{pi}',wording),
          'status':'SOURCE_CONFIRMED'})
        source_additions[sid]['appeal_categories']+=1
for sid,items in [('NEW-258',[('вода','WATER','ru','paragraph:5'),('мусор','WASTE','ru','paragraph:6'),
    ('дороги','ROAD','ru','paragraph:7'),('уличное освещение','STREET_LIGHT','ru','paragraph:8'),
    ('открытые люки','MANHOLE','ru','paragraph:9'),('земельные вопросы','LAND','ru','paragraph:10')]),
  ('NEW-264',[('Суу менен камсыздоо','WATER','ky','table:3:row:3'),
    ('Таштандыларды башкаруу','WASTE','ky','table:3:row:4'),
    ('Жолдор','ROAD','ky','table:3:row:5'),
    ('Көчө жарыктандыруусу','STREET_LIGHT','ky','table:3:row:6'),
    ('Ачык калган кудуктар','MANHOLE','ky','table:3:row:7')])]:
    for alias,topic,lang,loc in items:
        data['topic_aliases'].append({'topic_alias_id':ident('v4-topic-alias',sid,loc,alias),
          'canonical_topic':topic,'language':lang,'alias':alias,'source_ids':[sid],
          'provenance':provenance(sid,loc,alias),'effect':'LEXICAL_ONLY'})
        source_additions[sid]['topic_aliases']+=1

# Source relation and metrics are reproducible artifacts.
for name,rows in data.items():dump(name,rows)
metrics={'v3':original,'v4':{n:len(v) for n,v in data.items()},
 'by_source':{sid:dict(counts) for sid,counts in source_additions.items()},'review':review,
 'source_hashes':{sid:docs[sid]['hash'] for sid in docs}}
(ROOT/'build_metrics.json').write_text(json.dumps(metrics,ensure_ascii=False,indent=2),encoding='utf8')
print(json.dumps({'sources':len(data['sources']),'organizations':len(data['organizations']),
 'contacts':len(data['contacts']),'routes':len(data['routing_rules']),'chunks':len(data['rag_chunks'])},ensure_ascii=False))
