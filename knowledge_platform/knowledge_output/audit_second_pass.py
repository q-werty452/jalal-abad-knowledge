"""Read-only second-pass source inspection; outputs only audit_machine data."""
import hashlib
import io
import json
import re
import zipfile
from pathlib import Path

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / 'knowledge_output'
AUD = OUT / 'audit_machine'
DL = Path.home() / 'Downloads'
AUD.mkdir(exist_ok=True)

def read_jsonl(path):
    return [json.loads(x) for x in path.read_text(encoding='utf-8').splitlines() if x.strip()]

def write_jsonl(name, rows):
    with (AUD / name).open('w', encoding='utf-8') as f:
        for row in rows:
            f.write(json.dumps(row, ensure_ascii=False) + '\n')

sources = read_jsonl(OUT / 'machine/sources.jsonl')
existing_hash = {s['sha256']:s['source_id'] for s in sources if not s['duplicate_of']}
canonical_paths = list(OUT.glob('[0-1][0-9]_*.md')) + list((OUT/'machine').glob('*.jsonl'))
baseline = {str(p.relative_to(OUT)):hashlib.sha256(p.read_bytes()).hexdigest() for p in canonical_paths}
(AUD/'baseline_hashes.json').write_text(json.dumps(baseline, ensure_ascii=False, indent=2), encoding='utf-8')

def blocks(blob, ext):
    result=[]
    if ext=='.docx':
        doc=Document(io.BytesIO(blob))
        result += [{'locator':f'paragraph:{i}', 'text':p.text.strip()} for i,p in enumerate(doc.paragraphs,1) if p.text.strip()]
        for ti,table in enumerate(doc.tables,1):
            for ri,row in enumerate(table.rows,1):
                v=' | '.join(' '.join(c.text.split()) for c in row.cells)
                if v.strip(' |'): result.append({'locator':f'table:{ti}:row:{ri}', 'text':v})
    elif ext=='.xlsx':
        wb=load_workbook(io.BytesIO(blob),read_only=True,data_only=True)
        for ws in wb:
            for ri,row in enumerate(ws.iter_rows(values_only=True),1):
                v=' | '.join(str(c).replace('\n',' ') if c is not None else '' for c in row)
                if v.strip(' |'): result.append({'locator':f'sheet:{ws.title}:row:{ri}', 'text':v})
    elif ext=='.pdf':
        pdf=PdfReader(io.BytesIO(blob))
        for pi,page in enumerate(pdf.pages,1):
            for li,line in enumerate((page.extract_text() or '').splitlines(),1):
                if line.strip(): result.append({'locator':f'page:{pi}:line:{li}', 'text':' '.join(line.split())})
    elif ext in ('.txt','.md'):
        result=[{'locator':f'line:{i}', 'text':line.strip()} for i,line in enumerate(blob.decode('utf-8-sig',errors='replace').splitlines(),1) if line.strip()]
    return result

unrepresented=set(re.findall(r'SRC-\d+', (OUT/'16_FINAL_KNOWLEDGE_AUDIT.md').read_text(encoding='utf-8').split('## Source coverage after normalization')[1].split('## Recommended')[0]))
new_files=[
    'Жобо билим берүү институту (1).docx','Театрдын уставы (1) (1).docx',
    'Устав Филармония 2026-жыл июль (1).docx',
    'ЧАТ БОТ жөнүндө.УТРКЖалал-Абад.docxт туурасы (1).docx',
    'Положение о Таможне Джалал-Абад.docx',
    'Кыргыз темир жолу Түштүк бөлүмү филиалынын жобосу.docx',
    'Базар-Коргон финансы башкармалыгынын жобосу.docx',
    'ЖАМК №10 түзөтүү мекемеси жөнүндө жобо.docx']
items=[]
for name in new_files:
    p=DL/name
    b=p.read_bytes()
    items.append((name,b))
with zipfile.ZipFile(DL/'New Folder 1.zip') as z:
    for name in z.namelist():
        if name.lower().endswith('.docx'):
            items.append(('New Folder 1.zip::'+name,z.read(name)))
new=[]; extracted=[]
for i,(name,blob) in enumerate(items,231):
    sha=hashlib.sha256(blob).hexdigest()
    rec={'source_id':f'NEW-{i}', 'source_file':name, 'sha256':sha,
         'duplicate_of':existing_hash.get(sha),'status':'duplicate' if sha in existing_hash else 'new_unpromoted',
         'size':len(blob),'container_archive':'New Folder 1.zip' if '::' in name else None}
    new.append(rec)
    if not rec['duplicate_of']:
        try: extracted.append({'source_id':rec['source_id'],'source_file':name,'blocks':blocks(blob,'.docx')})
        except Exception as exc: rec['status']='blocked';rec['error']=str(exc)
write_jsonl('new_sources.jsonl',new)

for src in sources:
    if src['source_id'] not in unrepresented: continue
    try:
        if src['container_archive']:
            with zipfile.ZipFile(ROOT/src['container_archive']) as z: blob=z.read(src['archive_entry'])
        else: blob=(ROOT/src['source_file']).read_bytes()
        extracted.append({'source_id':src['source_id'],'source_file':src['source_file'],
                          'blocks':blocks(blob,'.'+src['file_type'].lower())})
    except Exception as exc:
        extracted.append({'source_id':src['source_id'],'source_file':src['source_file'],'blocks':[],'error':str(exc)})
write_jsonl('_review_extract.jsonl',extracted)
terms=re.compile(r'телефон|тел\.|моб\.|\+996|адрес|дарек|расположен|находится|при[её]м|кабыл ал|график|часы работы|иш убак|руководител|жетекчи|директор|начальник|башчы|школ|мектеп|образован|билим бер|вода|водоканал|суу|канализац|люк|дорог|жол|освещен|жарык|земл|жер|аренд|ижара|запис|жазыл|обращен|кайрыл|заявлен|арыз|айыл|район|шаар|город',re.I)
hits=[]
for rec in extracted:
    bs=rec['blocks']; matches=[b for b in bs if terms.search(b['text'])]
    hits.append({'source_id':rec['source_id'],'source_file':rec['source_file'],
                 'blocks':len(bs),'matches':len(matches),'sample':[{'locator':b['locator'],'text':b['text'][:260]} for b in matches[:12]]})
write_jsonl('_review_hits.jsonl',hits)
print(json.dumps({'unrepresented_expected':len(unrepresented),'reopened':sum(x['source_id'] in unrepresented for x in extracted),
                  'new_total':len(new),'new_unique':sum(not x['duplicate_of'] for x in new),
                  'review_blocks':sum(len(x['blocks']) for x in extracted)},ensure_ascii=False))
