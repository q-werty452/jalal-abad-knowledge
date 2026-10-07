"""Internal, repeatable source extraction. Never publish .extracted.jsonl directly."""
import hashlib
import io
import json
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "knowledge_output"
sources = [json.loads(x) for x in (OUT / "machine/sources.jsonl").read_text(encoding="utf-8").splitlines()]

def content(src):
    if src["container_archive"]:
        with ZipFile(ROOT / src["container_archive"]) as z:
            return z.read(src["archive_entry"])
    return (ROOT / src["source_file"]).read_bytes()

def extract_docx(blob):
    doc = Document(io.BytesIO(blob))
    blocks = []
    for n, p in enumerate(doc.paragraphs, 1):
        v = p.text.strip()
        if v:
            blocks.append({"locator": f"paragraph:{n}", "text": v})
    for ti, table in enumerate(doc.tables, 1):
        for ri, row in enumerate(table.rows, 1):
            vals = [" ".join(c.text.split()) for c in row.cells]
            if any(vals):
                blocks.append({"locator": f"table:{ti}:row:{ri}", "text": " | ".join(vals)})
    return blocks

def extract_xlsx(blob):
    wb = load_workbook(io.BytesIO(blob), read_only=True, data_only=True)
    blocks = []
    sheet_meta = []
    for ws in wb:
        first = []
        for ri, row in enumerate(ws.iter_rows(values_only=True), 1):
            vals = [str(v).replace("\n", " ").strip() if v is not None else "" for v in row]
            if any(vals):
                if len(first) < 1:
                    first.append({"row": ri, "values": vals[:30]})
                blocks.append({"locator": f"sheet:{ws.title}:row:{ri}", "text": " | ".join(vals)})
        sheet_meta.append({"sheet": ws.title, "rows": ws.max_row, "columns": ws.max_column, "first_nonempty_rows": first})
    return blocks, sheet_meta

def extract_pdf(blob):
    reader = PdfReader(io.BytesIO(blob))
    blocks = []
    for pn, page in enumerate(reader.pages, 1):
        try:
            txt = page.extract_text() or ""
        except Exception:
            txt = ""
        for j, line in enumerate(txt.splitlines(), 1):
            v = " ".join(line.split())
            if v:
                blocks.append({"locator": f"page:{pn}:line:{j}", "text": v})
    return blocks, len(reader.pages)

rows = []
for src in sources:
    rec = {"source_id": src["source_id"], "source_file": src["source_file"], "file_type": src["file_type"], "language": src["language"], "status": "duplicate" if src["duplicate_of"] else "processed", "duplicate_of": src["duplicate_of"], "blocks": [], "sheets": [], "page_count": None, "error": None}
    if src["duplicate_of"]:
        rows.append(rec)
        continue
    try:
        blob = content(src)
        typ = src["file_type"]
        if typ == "docx":
            rec["blocks"] = extract_docx(blob)
        elif typ == "xlsx":
            rec["blocks"], rec["sheets"] = extract_xlsx(blob)
        elif typ == "pdf":
            rec["blocks"], rec["page_count"] = extract_pdf(blob)
            raw = "".join(b["text"] for b in rec["blocks"])
            letters = sum(ch.isalpha() for ch in raw)
            cyrillic = sum("\u0400" <= ch <= "\u052f" for ch in raw)
            if letters > 500 and cyrillic / letters < 0.3:
                rec["status"] = "blocked"
                rec["error"] = "Extracted glyphs are unreadable; OCR or font-map repair required"
                rec["blocks"] = []
        elif typ in ("md", "txt"):
            rec["blocks"] = [{"locator": f"line:{i}", "text": line} for i, line in enumerate(blob.decode("utf-8-sig", errors="replace").splitlines(), 1) if line.strip()]
        else:
            rec["status"] = "blocked"
            rec["error"] = f"Unsupported format: {typ}"
        if not rec["blocks"] and rec["status"] == "processed":
            rec["status"] = "blocked"
            rec["error"] = "No extractable text; visual/OCR review required"
    except Exception as exc:
        rec["status"] = "blocked"
        rec["error"] = f"{type(exc).__name__}: {exc}"
    rows.append(rec)

with (OUT / ".extracted.jsonl").open("w", encoding="utf-8") as f:
    for row in rows:
        f.write(json.dumps(row, ensure_ascii=False) + "\n")

with (OUT / "machine/sources.jsonl").open("w", encoding="utf-8") as f:
    for src, row in zip(sources, rows):
        src["status"] = row["status"]
        src["extraction_error"] = row["error"]
        src["page_count"] = row["page_count"]
        src["sheet_structure"] = row["sheets"]
        src["block_count"] = len(row["blocks"])
        f.write(json.dumps(src, ensure_ascii=False) + "\n")

from collections import Counter
print(Counter(row["status"] for row in rows))
print("total blocks", sum(len(row["blocks"]) for row in rows))
