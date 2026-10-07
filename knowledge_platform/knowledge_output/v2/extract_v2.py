"""Read-only extraction of customer sources for the v2 promotion pass."""
import io
import json
import zipfile
from pathlib import Path

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader

ROOT=Path(__file__).resolve().parents[2]
OUT=ROOT/'knowledge_output'
V2=OUT/'v2'
V2.mkdir(exist_ok=True)
S=[json.loads(x) for x in (OUT/'machine/sources.jsonl').read_text(encoding='utf-8').splitlines()]
N=[json.loads(x) for x in (OUT/'audit_machine/new_sources.jsonl').read_text(encoding='utf-8').splitlines()]

def extract(blob,typ):
    result=[]
    if typ=='docx':
        doc=Document(io.BytesIO(blob))
        result.extend({'locator':f'paragraph:{i}','text':p.text.strip()} for i,p in enumerate(doc.paragraphs,1) if p.text.strip())
        for ti,table in enumerate(doc.tables,1):
            for ri,row in enumerate(table.rows,1):
                text=' | '.join(' '.join(c.text.split()) for c in row.cells)
                if text.strip(' |'):result.append({'locator':f'table:{ti}:row:{ri}','text':text})
    elif typ=='xlsx':
        wb=load_workbook(io.BytesIO(blob),read_only=True,data_only=True)
        for ws in wb:
            for ri,row in enumerate(ws.iter_rows(values_only=True),1):
                text=' | '.join(str(c).replace('\n',' ').strip() if c is not None else '' for c in row)
                if text.strip(' |'):result.append({'locator':f'sheet:{ws.title}:row:{ri}','text':text})
    elif typ=='pdf':
        pdf=PdfReader(io.BytesIO(blob))
        for pn,page in enumerate(pdf.pages,1):
            for ln,line in enumerate((page.extract_text() or '').splitlines(),1):
                if line.strip():result.append({'locator':f'page:{pn}:line:{ln}','text':' '.join(line.split())})
    elif typ in ('md','txt'):
        result.extend({'locator':f'line:{i}','text':x} for i,x in enumerate(blob.decode('utf-8-sig',errors='replace').splitlines(),1) if x.strip())
    return result

rows=[]
for s in S:
    if s['status']!='processed':continue
    try:
        if s.get('container_archive'):
            with zipfile.ZipFile(ROOT/s['container_archive']) as z:blob=z.read(s['archive_entry'])
        else:blob=(ROOT/s['source_file']).read_bytes()
        blocks=extract(blob,s['file_type'].lower())
        rows.append({'source_id':s['source_id'],'source_file':s['source_file'],'blocks':blocks})
    except Exception as exc:
        rows.append({'source_id':s['source_id'],'source_file':s['source_file'],'blocks':[],'error':repr(exc)})
for s in N:
    if s['duplicate_of']:continue
    try:
        if s['container_archive']:
            with zipfile.ZipFile(Path.home()/'Downloads'/'New Folder 1.zip') as z:blob=z.read(s['source_file'].split('::',1)[1])
        else:blob=(Path.home()/'Downloads'/s['source_file']).read_bytes()
        rows.append({'source_id':s['source_id'],'source_file':s['source_file'],'blocks':extract(blob,'docx')})
    except Exception as exc:
        rows.append({'source_id':s['source_id'],'source_file':s['source_file'],'blocks':[],'error':repr(exc)})
with (V2/'.source_blocks.jsonl').open('w',encoding='utf-8') as f:
    for r in rows:f.write(json.dumps(r,ensure_ascii=False)+'\n')
print(json.dumps({'sources':len(rows),'blocks':sum(len(r['blocks']) for r in rows),'errors':[(r['source_id'],r['error']) for r in rows if r.get('error')]},ensure_ascii=False))
